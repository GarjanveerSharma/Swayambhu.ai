"""
Alag-alag file types se text nikalna.
Har loader page-wise text ki list deta hai: ["page 1 ka text", "page 2 ka text", ...]
Page number baad me sources dikhane ke kaam aayega.
"""
from pathlib import Path

DOCX_CHARS_PER_PAGE = 3000  # DOCX me asli pages nahi hote, isliye approx tukde


class ExtractionError(Exception):
    """User ko dikhane layak error."""


def _load_pdf(path: Path) -> list[str]:
    import fitz  # pymupdf

    try:
        with fitz.open(path) as doc:
            pages = [page.get_text("text") for page in doc]
    except Exception as exc:
        raise ExtractionError("PDF khul nahi paya. File kharab ya password-protected ho sakti hai.") from exc

    if not any(p.strip() for p in pages):
        # Scanned PDF me text layer nahi hoti
        raise ExtractionError("Is PDF me text nahi mila (scanned lagta hai). OCR support agle part me aayega.")
    return pages


def _load_docx(path: Path) -> list[str]:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError("DOCX file khul nahi payi.") from exc

    lines = [p.text for p in d.paragraphs if p.text.strip()]
    # Tables ka text bhi lo
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return _split_into_pages("\n".join(lines))


def _load_txt(path: Path) -> list[str]:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            return _split_into_pages(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Text file ki encoding samajh nahi aayi.")


def _split_into_pages(text: str) -> list[str]:
    text = text.strip()
    if not text:
        raise ExtractionError("File khali hai.")
    pages, current, size = [], [], 0
    for line in text.splitlines():
        current.append(line)
        size += len(line) + 1
        if size >= DOCX_CHARS_PER_PAGE:
            pages.append("\n".join(current))
            current, size = [], 0
    if current:
        pages.append("\n".join(current))
    return pages


LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
}

# Ye types upload ho sakte hain, par abhi text nahi nikal sakte (OCR agle part me)
OCR_ONLY = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}

ALLOWED_EXTENSIONS = set(LOADERS) | OCR_ONLY


def extract_pages(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext in OCR_ONLY:
        raise ExtractionError("Image/scan se text nikalna (OCR) agle part me aayega.")
    loader = LOADERS.get(ext)
    if not loader:
        raise ExtractionError(f"'{ext}' file type support nahi hai.")
    return loader(path)